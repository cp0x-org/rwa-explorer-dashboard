import copy
import hashlib
import json
from pathlib import Path

import pytest

from rwa_pipeline import cmc, collect, portals
from rwa_pipeline.analyze import build, find_relations, normalize_opportunity, parent_index
from rwa_pipeline.api import ApiError
from rwa_pipeline.config import Asset, isin_is_valid, load_assets, load_settings, write_json
from rwa_pipeline.export import build_export

FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "nvda_snapshot.json").read_text())
SETTINGS = load_settings()
NVDAC = "base:0xb20000000000000000000078ee7ce2fe4908108c"
NVDA = Asset("NVDA", "NVIDIA", "stock", "US67066G1040", True)
AAPL = Asset("AAPL", "Apple", "stock", "US0378331005", True)


def snapshot():
    return copy.deepcopy(FIXTURE)


def by_symbol(export, symbol):
    return next(t for t in export["tokens"] if t["symbol"] == symbol)


def record(key, tokens, platform="beefy", tags=("#yield",), liquidity=50000, apy="5", symbol="X"):
    network, _, address = key.partition(":")
    return {
        "key": key,
        "network": network,
        "address": address,
        "platform": platform,
        "symbol": symbol,
        "tokens": tokens,
        "liquidity": liquidity,
        "metrics": {"apy": apy},
        "metadata": {"tags": list(tags)},
        "isPortable": True,
        "updatedAt": FIXTURE["fetched_at"],
    }


# ---------- config / CMC ----------


def test_isin_check_digit():
    assert all(isin_is_valid(a.isin) for a in load_assets())
    assert not isin_is_valid("US67066G1041")
    assert not isin_is_valid("US67066G104")


def test_deployments_per_network_and_address_case():
    deps = cmc.parse_deployments(FIXTURE["cmc"]["info"]["36992"], SETTINGS["chain_mapping"])  # NVDAX
    keys = {d["key"] for d in deps if d["key"]}
    # same EVM address on two networks -> two deployments
    assert {"ethereum:0xc845b2894dbddd03858fd2d643b4ef725fe0849d", "bsc:0xc845b2894dbddd03858fd2d643b4ef725fe0849d"} <= keys
    solana = next(d for d in deps if d["cmc_network"] == "Solana")
    assert solana["address"] == "Xsc9qvGR1efVDFGLrVsmkzv3qi45LTBjeUKSPmx9qEh"  # non-EVM case kept
    assert solana["status"] == "unsupported_network"

    info = {
        "contract_address": [
            {"contract_address": "0xB989ad9b91886b1Aaed8DaADb26F028b29b40945", "platform": {"name": "HyperEVM", "coin": {"id": "32196"}}},
            {"contract_address": "0xabc", "platform": {"name": "Hyperliquid", "coin": {"id": "32196"}}},
            {"contract_address": "0xdef", "platform": {"name": "Some New Chain", "coin": {"id": "1"}}},
        ]
    }
    statuses = {d["cmc_network"]: (d["network"], d["status"]) for d in cmc.parse_deployments(info, SETTINGS["chain_mapping"])}
    assert statuses["HyperEVM"] == ("hyperevm", None)  # shares coin id with Hyperliquid: mapped by name
    assert statuses["Hyperliquid"] == (None, "unsupported_network")
    assert statuses["Some New Chain"] == (None, "unmapped_network")


class FakeApi:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.credits = 0

    def get(self, path, params):
        self.calls.append((path, params))
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


def test_cmc_parsing_shapes():
    api = FakeApi([{"status": {"error_code": "0", "credit_count": 1}, "data": {"rwa_assets": [FIXTURE["cmc"]["rwa"]]}}])
    quotes = cmc.fetch_rwa_quotes(api, ["NVDA"])
    rwa, err = cmc.pick_rwa_asset(quotes["NVDA"], NVDA, None)
    assert err is None and {t["crypto_id"] for t in rwa["tokens"]} >= {42461, 38093}

    api = FakeApi([{"status": {"error_code": 1002, "error_message": "bad key"}}])
    with pytest.raises(ApiError):
        cmc.fetch_rwa_quotes(api, ["NVDA"])

    # one invalid id fails the batch -> retried without it
    api = FakeApi(
        [
            ApiError("cmc", 400, "Invalid value for 'id': '39839'"),
            {"status": {"error_code": 0}, "data": {"42461": FIXTURE["cmc"]["info"]["42461"]}},
        ]
    )
    info, invalid = cmc.fetch_crypto_info(api, [42461, 39839])
    assert invalid == [39839] and 42461 in info
    assert api.calls[1][1]["id"] == "42461"

    api = FakeApi([{"status": {"error_code": "0"}, "data": {"name": "Coinbase", "tokens": [{"crypto_id": 1}], "num_tokens": 1}}])
    assert cmc.fetch_issuer(api, "abc")["name"] == "Coinbase"


def test_portals_pagination_until_more_false():
    pages = [
        {"tokens": [{"key": "base:0x1"}, {"key": "base:0x2"}], "more": True, "totalItems": 3},
        {"tokens": [{"key": "base:0x2"}, {"key": "base:0x3"}], "more": False, "totalItems": 3},
    ]
    api = FakeApi(pages)
    items, meta = portals.search(api, "NVDA", 2)
    assert set(items) == {"base:0x1", "base:0x2", "base:0x3"}
    assert [c[1]["page"] for c in api.calls] == [0, 1] and all(c[1]["limit"] == 2 for c in api.calls)
    assert meta["complete"] and meta["pages"] == 2

    api = FakeApi([{"tokens": [{"key": "base:0x1"}], "more": True, "totalItems": 999}] * portals.MAX_PAGES)
    _, meta = portals.search(api, "X", 1)
    assert not meta["complete"]

    assert portals.minimal_terms(["NVDA", "NVDAc", "NVDAon", "rNVDA", "nvda"]) == ["NVDA"]


# ---------- graph / classification ----------


def test_fixture_graph_root_lp_vault_and_root_exclusion():
    export = build([snapshot()], {}, SETTINGS, {})
    nvdac = by_symbol(export, "NVDAc")
    rel = {r["key"]: r for r in nvdac["relations"]}
    opps = {o["key"]: o for o in export["opportunities"]}

    assert NVDAC not in rel  # the issuer root (coinbase-tokenized-stocks platform) is not an opportunity
    reward = next(k for k in rel if opps[k]["symbol"] == "RCOWAERODROMEBASEUSDC-NVDAC")
    assert rel[reward]["depth"] == 2
    assert opps[rel[reward]["path"][1]]["symbol"] == "COWAERODROMEBASEUSDC-NVDAC"
    # Spectra market reachable via PT and via fNVDAc: counted once with min depth
    market = [k for k in rel if opps[k]["symbol"] == "PT-FNVDAC(NVDAC)-2026/12/28/FNVDAC"]
    assert len(market) == 1 and rel[market[0]]["depth"] == 2
    # Morpho market lists only USDC in tokens[]: symbol match only -> unconfirmed, not counted
    morpho = [k for k in opps if opps[k]["platform"] == "morpho"]
    assert morpho and all(k not in rel for k in morpho) and set(morpho) <= set(nvdac["unconfirmed_keys"])
    assert nvdac["usefulness_score"] == nvdac["counts"]["eligible_opportunities_count"] > 0

    # wrappers found in composition are kept as raw but never eligible
    nvdax = by_symbol(export, "NVDAX")
    wrappers = [r["key"] for r in nvdax["relations"] if opps[r["key"]]["category"] == "wrapper"]
    assert wrappers and all(opps[k]["eligibility"] == "category_excluded" for k in wrappers)

    derivative = next(t for t in export["tokens"] if t["kind"] == "derivative")
    assert derivative["id"] not in export["underlyings"][0]["token_ids"]


def test_cycle_max_depth_and_swap_paths():
    root = "base:0xroot"
    records = {
        "base:0xa": record("base:0xa", ["0xroot", "0xb"]),
        "base:0xb": record("base:0xb", ["0xa"]),  # cycle a <-> b
        "base:0xc": record("base:0xc", ["0xb"]),
        "base:0xd": record("base:0xd", ["0xc"]),  # depth 4
        "base:0xusdcvault": record("base:0xusdcvault", ["0xusdc"]),  # sell root for USDC -> vault: no edge
    }
    found = find_relations([root], parent_index(records), set(), max_depth=3)
    assert {k: v["depth"] for k, v in found.items()} == {"base:0xa": 1, "base:0xb": 2, "base:0xc": 3}


def test_symbol_match_with_other_address_is_not_confirmed():
    snap = snapshot()
    fake = record("base:0xfeed", ["0x0000000000000000000000000000000000000bad"], platform="aerodrome", tags=("#clp",), symbol="NVDAC/USDC")
    snap["portals"]["records"][fake["key"]] = fake
    nvdac = by_symbol(build([snap], {}, SETTINGS, {}), "NVDAc")
    assert fake["key"] not in {r["key"] for r in nvdac["relations"]}
    assert fake["key"] in nvdac["unconfirmed_keys"]


def test_v4_pool_id_non_portable_lp_stays_eligible():
    v4 = next(r for r in FIXTURE["portals"]["records"].values() if r["platform"] == "uniswap-v4" and r["network"] == "base")
    opp = normalize_opportunity(v4, SETTINGS, {})
    assert opp["position_kind"] == "pool_id" and opp["receipt_token_key"] is None
    assert v4["isPortable"] is False and "sentinel_price" in opp["flags"]
    assert opp["category"] == "dex_lp" and opp["liquidity_semantics"] == "pool_tvl"


def test_metric_edge_cases():
    base = record("base:0x1", ["0xroot"], platform="morpho", tags=("#lending",))
    base["metrics"] = {"borrowApy": "5.1", "volumeUsd7d": "-120.5"}
    opp = normalize_opportunity(base, SETTINGS, {})
    assert opp["apy_percent"] is None and opp["borrow_apy_percent"] == "5.1"  # borrow cost is not yield
    assert "apy_missing" in opp["flags"] and "negative_volume_net_flow" in opp["flags"]
    assert opp["volume_usd_7d"] == "-120.5"  # no abs()

    weird = record("base:0x2", ["0xroot"], apy="25000")
    assert "apy_anomaly" in normalize_opportunity(weird, SETTINGS, {})["flags"]

    root = "base:0xroot"
    snap = snapshot()
    snap["portals"]["records"] = {
        "base:0x2": weird,
        "base:0x3": record("base:0x3", ["0xroot"], apy=None),
        "base:0x4": record("base:0x4", ["0xroot"], apy="7.5"),
    }
    snap["portals"]["exact"] = {root: record(root, [], platform="basic", tags=())}
    snap["portals"]["looked_up"] = [root]
    snap["cmc"]["info"]["42461"]["contract_address"] = [{"contract_address": "0xroot", "platform": {"name": "Base"}}]
    y = by_symbol(build([snap], {}, SETTINGS, {}), "NVDAc")["yield"]
    assert y["max_reported_apy_percent"] == "7.5" and y["yield_opportunities_count"] == 1
    assert y["anomaly_count"] == 1 and y["missing_apy_count"] == 1


def test_issuer_counts_shared_pool_once():
    snap = snapshot()
    ondo = next(t for t in snap["cmc"]["rwa"]["tokens"] if t["crypto_id"] == 38093)
    nvdac = next(t for t in snap["cmc"]["rwa"]["tokens"] if t["crypto_id"] == 42461)
    ondo["issuer_id"] = nvdac["issuer_id"]  # pretend both versions belong to one issuer
    snap["cmc"]["info"]["38093"]["contract_address"].append({"contract_address": "0xb20000000000000000000078ee7ce2fe4908108c", "platform": {"name": "Base"}})
    export = build([snap], {}, SETTINGS, {})
    issuer = next(i for i in export["issuers"] if i["id"] == nvdac["issuer_id"])
    own = [t for t in export["tokens"] if t["issuer_id"] == nvdac["issuer_id"]]
    union = {r["key"] for t in own for r in t["relations"] if next(o for o in export["opportunities"] if o["key"] == r["key"])["eligibility"] == "eligible"}
    assert issuer["issuer_usefulness_score"] == len(union) < sum(t["counts"]["eligible_opportunities_count"] for t in own)
    # market cap summed once per token, never per deployment
    from decimal import Decimal

    assert Decimal(issuer["basket_market_cap_usd"]) == sum(Decimal(str(t["market_cap"])) for t in (ondo, nvdac))


# ---------- sync modes ----------


def _patch_network(monkeypatch, quotes_calls, fail_search=False):
    monkeypatch.setattr(collect, "load_api_keys", lambda: ("k1", "k2"))
    aapl_rwa = {"rwa_id": 3, "symbol": "AAPL", "asset_type": "stock", "tokens": []}

    def quotes(api, tickers):
        quotes_calls.append(list(tickers))
        return {"NVDA": [FIXTURE["cmc"]["rwa"]], "AAPL": [aapl_rwa]}

    def search(api, term, limit):
        if fail_search:
            raise ApiError("portals", 503, "down")
        return {}, {"complete": True}

    monkeypatch.setattr(cmc, "fetch_rwa_quotes", quotes)
    monkeypatch.setattr(cmc, "fetch_crypto_info", lambda api, ids: ({int(k): v for k, v in FIXTURE["cmc"]["info"].items()}, []))
    monkeypatch.setattr(cmc, "fetch_issuer", lambda api, i: {"issuer_id": i, "name": "x"})
    monkeypatch.setattr(cmc, "fetch_rwa_info", lambda api, ids: {str(i): {"rwa_id": i, "website": None, "logo": None} for i in ids})
    monkeypatch.setattr(portals, "fetch_networks", lambda api: FIXTURE["portals"]["networks"])
    monkeypatch.setattr(portals, "fetch_platform_names", lambda api: {})
    monkeypatch.setattr(portals, "lookup", lambda api, keys, limit: {})
    monkeypatch.setattr(portals, "search", search)


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_add_is_noop_and_new_ticker_keeps_old_snapshot(tmp_path, monkeypatch):
    nvda_path = collect.snapshot_path(tmp_path, NVDA)
    write_json(nvda_path, snapshot())
    before = _digest(nvda_path)

    monkeypatch.setattr(collect, "load_api_keys", lambda: pytest.fail("no network expected"))
    assert collect.sync([NVDA], SETTINGS, tmp_path, None, refresh=False) == (0, None)

    calls = []
    _patch_network(monkeypatch, calls)
    code, report = collect.sync([NVDA, AAPL], SETTINGS, tmp_path, None, refresh=False)
    assert code == 0 and calls == [["AAPL"]]
    assert _digest(nvda_path) == before
    assert json.loads(collect.snapshot_path(tmp_path, AAPL).read_text())["status"] == "no_tracked_tokens"
    # a ticker resolved with tokens=[] counts as done: the next add is a no-op
    assert collect.sync([NVDA, AAPL], SETTINGS, tmp_path, None, refresh=False) == (0, None)


def test_refresh_is_scoped_and_failure_keeps_last_snapshot(tmp_path, monkeypatch):
    nvda_path, aapl_path = collect.snapshot_path(tmp_path, NVDA), collect.snapshot_path(tmp_path, AAPL)
    write_json(nvda_path, snapshot())
    write_json(aapl_path, {"underlying_id": AAPL.id, "status": "no_tracked_tokens"})
    nvda_before, aapl_before = _digest(nvda_path), _digest(aapl_path)

    calls = []
    _patch_network(monkeypatch, calls, fail_search=True)
    code, report = collect.sync([NVDA, AAPL], SETTINGS, tmp_path, ["NVDA"], refresh=True)
    assert code == 2 and "NVDA" in report["failed"]
    assert _digest(nvda_path) == nvda_before  # failed Portals search does not wipe positions
    assert _digest(aapl_path) == aapl_before

    _patch_network(monkeypatch, calls)
    code, _ = collect.sync([NVDA, AAPL], SETTINGS, tmp_path, ["NVDA"], refresh=True)
    assert code == 0 and calls[-1] == ["NVDA"] and _digest(aapl_path) == aapl_before

    with pytest.raises(Exception):
        collect.sync([NVDA, AAPL], SETTINGS, tmp_path, ["MSFT"], refresh=True)


# ---------- export ----------


def test_export_has_no_secrets_and_keeps_null_vs_zero(tmp_path, monkeypatch):
    monkeypatch.setenv("CMC_API_KEY", "secret-cmc-value")
    monkeypatch.setenv("PORTALS_API_KEY", "secret-portals-value")
    snap = snapshot()
    nvdac = next(t for t in snap["cmc"]["rwa"]["tokens"] if t["crypto_id"] == 42461)
    nvdac["price"], nvdac["volume_24h"] = None, 0
    write_json(collect.snapshot_path(tmp_path, NVDA), snap)

    export = build_export([NVDA, Asset("AAPL", "Apple", "stock", "US0378331005", False)], SETTINGS, tmp_path)
    text = json.dumps(export)
    assert "secret-cmc-value" not in text and "secret-portals-value" not in text and "Bearer" not in text
    token = by_symbol(export, "NVDAc")
    assert token["price_usd"] is None and token["volume_24h_usd"] == "0"
    assert [u["ticker"] for u in export["underlyings"]] == ["NVDA"]  # disabled assets are not in the active basket
    assert "NaN" not in text and "Infinity" not in text


# ---------- logos ----------


def test_website_icon_prefers_biggest_declared_size():
    import httpx

    from rwa_pipeline.logos import website_icon_url

    html = """<head>
      <link href="/fav_32.png" rel="shortcut icon"/>
      <link href="/touch_60.png" rel="apple-touch-icon"/>
      <link href="/touch_120.png" rel="apple-touch-icon" sizes="120x120"/>
    </head>"""
    client = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, text=html)))
    assert website_icon_url(client, "https://issuer.example/en/") == "https://issuer.example/touch_120.png"


def test_export_uses_local_logos_only(tmp_path):
    write_json(collect.snapshot_path(tmp_path, NVDA), snapshot())
    (tmp_path / "logos" / "tokens").mkdir(parents=True)
    (tmp_path / "logos" / "tokens" / "42461.png").write_bytes(b"\x89PNG fake")
    (tmp_path / "logos" / "underlyings").mkdir(parents=True)
    (tmp_path / "logos" / "underlyings" / "US67066G1040.png").write_bytes(b"\x89PNG nvda")
    write_json(tmp_path / "rwa_info.json", {"2": {"rwa_id": 2, "website": "https://www.nvidia.com", "logo": None}})
    frontend = tmp_path / "frontend"

    from rwa_pipeline.export import write_export

    export = build_export([NVDA], SETTINGS, tmp_path)
    write_export(export, tmp_path, frontend)
    assert by_symbol(export, "NVDAc")["logo"] == "logos/tokens/42461.png"
    assert by_symbol(export, "NVDAon")["logo"] is None  # no file -> no remote URL fallback
    assert (frontend / "public" / "logos" / "tokens" / "42461.png").read_bytes() == b"\x89PNG fake"
    assert "coinmarketcap.com/static" not in json.dumps(export)
    underlying = export["underlyings"][0]
    assert underlying["logo"] == "logos/underlyings/US67066G1040.png" and underlying["website"] == "https://www.nvidia.com"
    assert (frontend / "public" / "logos" / "underlyings" / "US67066G1040.png").exists()


# ---------- derived metrics and tables ----------


def test_aggregate_yield_and_defi_presence_metrics():
    from decimal import Decimal

    export = build([snapshot()], {}, SETTINGS, {"beefy": "Beefy"})
    opps = {o["key"]: o for o in export["opportunities"]}
    nvdac = by_symbol(export, "NVDAc")
    reported = [
        Decimal(opps[r["key"]]["apy_percent"])
        for r in nvdac["relations"]
        if opps[r["key"]]["eligibility"] == "eligible"
        and opps[r["key"]]["apy_percent"] is not None
        and Decimal(opps[r["key"]]["apy_percent"]) > 0
        and "apy_anomaly" not in opps[r["key"]]["flags"]
    ]
    assert reported and nvdac["yield"]["yield_opportunities_count"] == len(reported)
    assert Decimal(nvdac["yield"]["average_reported_apy_percent"]) == (sum(reported) / len(reported)).quantize(Decimal("0.0001"))
    assert Decimal(nvdac["yield"]["max_reported_apy_percent"]) == max(reported)

    underlying = export["underlyings"][0]
    included = [t for t in export["tokens"] if t["id"] in underlying["token_ids"]]
    assert underlying["tokens_in_defi_count"] == sum(1 for t in included if t["counts"]["eligible_opportunities_count"] > 0) > 0
    assert underlying["eligible_platforms"] and len(underlying["eligible_platforms"]) == underlying["eligible_protocols_count"]
    assert export["summary"]["tokens_in_defi_count"] == underlying["tokens_in_defi_count"]
    assert export["summary"]["average_reported_apy_percent"] is not None

    issuer = next(i for i in export["issuers"] if i["id"] == nvdac["issuer_id"])
    assert issuer["tokens_in_defi_count"] <= issuer["basket_tokens_count"]
    assert len(issuer["eligible_platforms"]) == issuer["eligible_protocols_count"]
    assert issuer["average_reported_apy_percent"] is not None and issuer["max_reported_apy_percent"] is not None


def test_tables_cover_every_token_relation_and_deployment(tmp_path):
    import csv

    from openpyxl import load_workbook

    from rwa_pipeline.tables import build_tables, write_tables

    snap = snapshot()
    nvdac_raw = next(t for t in snap["cmc"]["rwa"]["tokens"] if t["crypto_id"] == 42461)
    nvdac_raw["price"], nvdac_raw["volume_24h"] = None, 0
    write_json(collect.snapshot_path(tmp_path, NVDA), snap)
    export = build_export([NVDA], SETTINGS, tmp_path)
    tables = {t.name: t for t in build_tables(export)}

    assert set(tables) == {"summary", "underlyings", "tokens", "issuers", "opportunities", "relations", "deployments", "breakdowns"}
    assert len(tables["tokens"].rows) == len(export["tokens"])  # excluded tokens are kept, marked by kind
    assert len(tables["underlyings"].rows) == 1 and len(tables["opportunities"].rows) == len(export["opportunities"])
    assert len(tables["deployments"].rows) == sum(len(t["deployments"]) for t in export["tokens"])
    assert len(tables["relations"].rows) == sum(len(t["relations"]) + len(t["unconfirmed_keys"]) for t in export["tokens"])
    links = {r["link"] for r in tables["relations"].rows}
    assert links == {"direct", "indirect", "unconfirmed"}
    assert all(r["status"] == "unconfirmed" and r["depth"] is None for r in tables["relations"].rows if r["link"] == "unconfirmed")
    for t in tables.values():
        assert all(set(row) <= set(t.columns) for row in t.rows), t.name
    nvdac = next(r for r in tables["tokens"].rows if r["symbol"] == "NVDAc")
    assert nvdac["cmc_url"].startswith("https://coinmarketcap.com/currencies/") and nvdac["eligible_platforms"]
    assert {m["metric"] for m in tables["summary"].rows} >= {"generated_at", "setting.min_eligible_liquidity_usd", "eligible_opportunities_count"}

    paths = write_tables(export, "all", None, tmp_path)
    csv_paths = [p for p in paths if p.suffix == ".csv"]
    assert {p.stem for p in csv_paths} == set(tables) and all(p.parent == tmp_path / "exports" / "csv" for p in csv_paths)
    with (tmp_path / "exports" / "csv" / "tokens.csv").open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == len(export["tokens"])
    row = next(r for r in rows if r["symbol"] == "NVDAc")
    assert row["market_cap_usd"] == by_symbol(export, "NVDAc")["market_cap_usd"]  # decimal string, no exponent
    assert row["price_usd"] == "" and row["volume_24h_usd"] == "0"  # null -> empty cell, zero stays zero

    xlsx = tmp_path / "exports" / "rwa.xlsx"
    assert xlsx in paths
    wb = load_workbook(xlsx, read_only=True)
    assert wb.sheetnames == ["README"] + list(tables)
    header, *body = list(wb["tokens"].iter_rows(values_only=True))
    assert list(header) == tables["tokens"].columns and len(body) == len(export["tokens"])
    cell = dict(zip(header, next(r for r in body if r[header.index("symbol")] == "NVDAc")))
    assert isinstance(cell["market_cap_usd"], float) and cell["fetched_at"].year >= 2025  # numbers and dates are typed
    assert cell["price_usd"] is None and cell["volume_24h_usd"] == 0


def test_tables_out_overrides_default_target(tmp_path):
    from rwa_pipeline.tables import write_tables

    write_json(collect.snapshot_path(tmp_path, NVDA), snapshot())
    export = build_export([NVDA], SETTINGS, tmp_path)
    (xlsx,) = write_tables(export, "xlsx", tmp_path / "custom" / "basket.xlsx", tmp_path)
    assert xlsx == tmp_path / "custom" / "basket.xlsx" and xlsx.exists()
    csvs = write_tables(export, "csv", tmp_path / "csv-out", tmp_path)
    assert csvs and all(p.parent == tmp_path / "csv-out" for p in csvs)
