"""Flat tables for manual analysis: the export (`rwa.json`) as CSV files or one XLSX workbook.

Every table is a list of columns plus rows of plain values. Nested parts of the export (deployments, relations,
breakdowns) become their own long-format tables, so each sheet pivots in Excel/pandas without any parsing.
Decimal strings of the export stay numbers (float), `null` stays an empty cell, lists are joined with "; ".
"""

import csv
import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

LIST_SEPARATOR = "; "
CMC_URL = "https://coinmarketcap.com/currencies/{slug}/"

COUNT_FIELDS = [
    "raw_confirmed_opportunities_count",
    "eligible_opportunities_count",
    "unknown_liquidity_count",
    "direct_eligible_count",
    "indirect_eligible_count",
    "eligible_protocols_count",
    "eligible_networks_count",
]
YIELD_FIELDS = ["yield_opportunities_count", "max_reported_apy_percent", "average_reported_apy_percent", "median_reported_apy_percent"]
OPPORTUNITY_FIELDS = [
    "key",
    "symbol",
    "name",
    "platform",
    "platform_name",
    "network",
    "category",
    "position_kind",
    "address",
    "receipt_token_key",
    "is_portable",
    "liquidity_usd",
    "liquidity_semantics",
    "apy_percent",
    "base_apy_percent",
    "reward_apy_percent",
    "borrow_apy_percent",
    "volume_usd_1d",
    "volume_usd_7d",
    "eligibility",
    "flags",
    "tags",
    "updated_at",
]
# Columns whose values are decimal strings in the export; everything else is kept as is.
NUMERIC_SUFFIXES = ("_usd", "_percent", "_score", "_count", "per_underlying", "per_token", "_size")
TIME_FIELDS = {"fetched_at", "oldest_fetched_at", "newest_fetched_at", "updated_at", "generated_at"}


@dataclass
class Table:
    name: str
    columns: list[str]
    rows: list[dict]
    description: str = ""


def _join(values) -> str | None:
    if not values:
        return None
    return LIST_SEPARATOR.join(str(v) for v in values)


def _as_number(value):
    """Decimal strings become floats so spreadsheets sort and sum them; anything else is returned untouched."""
    if isinstance(value, str):
        try:
            d = Decimal(value)
        except InvalidOperation:
            return value
        return float(d) if d.is_finite() else value
    return value


def _pick(item: dict, fields: list[str]) -> dict:
    return {f: item.get(f) for f in fields}


def build_tables(export: dict) -> list[Table]:
    underlying_by_id = {u["id"]: u for u in export["underlyings"]}
    issuer_by_id = {i["id"]: i for i in export["issuers"]}
    opp_by_key = {o["key"]: o for o in export["opportunities"]}
    token_by_id = {t["id"]: t for t in export["tokens"]}
    basket = [t for t in export["tokens"] if t["excluded_reason"] is None]

    # ---- summary: one metric per row, settings and provenance included ----
    summary_rows = [
        {"metric": "generated_at", "value": export["generated_at"]},
        {"metric": "schema_version", "value": export["schema_version"]},
        {"metric": "score_version", "value": export["score_version"]},
        *({"metric": f"setting.{k}", "value": v} for k, v in export["settings"].items()),
        *({"metric": k, "value": v} for k, v in export["summary"].items()),
    ]

    # ---- underlyings ----
    underlying_rows = []
    for u in export["underlyings"]:
        row = _pick(
            u,
            [
                "id", "ticker", "name", "isin", "asset_type", "status", "cmc_rwa_id", "cmc_name", "website",
                "tokenized_market_cap_usd", "tokenized_volume_24h_usd", "average_tokenized_price_usd",
                "versions_count", "issuers_count", "tokens_in_defi_count",
            ],
        )
        row["excluded_tokens_count"] = len(u["excluded_token_ids"])
        row.update(_pick(u, COUNT_FIELDS))
        row.update(_pick(u, YIELD_FIELDS))
        row["eligible_platforms"] = _join(u.get("eligible_platforms"))
        row["tokens"] = _join(token_by_id[i]["symbol"] for i in u["token_ids"] if i in token_by_id)
        row["issuers"] = _join(sorted({token_by_id[i]["issuer_name"] for i in u["token_ids"] if i in token_by_id and token_by_id[i]["issuer_name"]}))
        row["coverage"] = u["coverage"]
        row["search_terms"] = _join(u["search_terms"])
        row["search_complete"] = u["search_complete"]
        row["fetched_at"] = u["fetched_at"]
        row["logo"] = u["logo"]
        underlying_rows.append(row)

    # ---- tokens (all, excluded ones marked by kind/excluded_reason) ----
    token_rows, deployment_rows, relation_rows, breakdown_rows = [], [], [], []
    for t in export["tokens"]:
        u = underlying_by_id.get(t["underlying_id"]) or {}
        y = t["yield"]
        row = _pick(t, ["id", "symbol", "name", "kind", "excluded_reason", "ticker"])
        row["underlying_name"] = u.get("name")
        row["underlying_id"] = t["underlying_id"]
        row.update(_pick(t, ["issuer_id", "issuer_name"]))
        row["rank_in_underlying"] = t.get("rank_in_underlying")
        row.update(_pick(t, ["price_usd", "market_cap_usd", "volume_24h_usd", "usefulness_score"]))
        row.update(_pick(t["counts"], COUNT_FIELDS + ["unconfirmed_candidates_count"]))
        row["deployments_count"] = len(t["deployments"])
        row["indexed_deployments_count"] = sum(1 for d in t["deployments"] if d["status"] == "indexed")
        row["deployment_networks"] = _join(sorted({d["cmc_network"] for d in t["deployments"]}))
        row["eligible_networks"] = _join(k for k, v in t["by_network"].items() if v["eligible"])
        row["eligible_platforms"] = _join(sorted({opp_by_key[r["key"]]["platform_name"] for r in t["relations"] if opp_by_key[r["key"]]["eligibility"] == "eligible"}))
        row["eligible_categories"] = _join(k for k, v in t["by_category"].items() if v["eligible"])
        row["yield_status"] = y["status"]
        row.update(_pick(y, YIELD_FIELDS + ["max_reported_apy_key", "missing_apy_count", "anomaly_count"]))
        row["coverage"] = t["coverage"]
        row["fetched_at"] = t["fetched_at"]
        row["cmc_slug"] = t["cmc_slug"]
        row["cmc_url"] = CMC_URL.format(slug=t["cmc_slug"]) if t["cmc_slug"] else None
        row["logo"] = t["logo"]
        token_rows.append(row)

        for d in t["deployments"]:
            deployment_rows.append(
                {
                    "token_id": t["id"],
                    "token_symbol": t["symbol"],
                    "ticker": t["ticker"],
                    "issuer_name": t["issuer_name"],
                    "token_kind": t["kind"],
                    **_pick(d, ["cmc_network", "network", "address", "key", "status", "portals_symbol", "decimals", "portals_platform"]),
                }
            )

        token_cols = {"token_id": t["id"], "token_symbol": t["symbol"], "ticker": t["ticker"], "issuer_name": t["issuer_name"], "token_kind": t["kind"]}
        for r in t["relations"]:
            o = opp_by_key[r["key"]]
            relation_rows.append(
                {
                    **token_cols,
                    "link": "direct" if r["depth"] == 1 else "indirect",
                    "depth": r["depth"],
                    "path": _join(r["path"]),
                    "status": o["eligibility"],
                    **_opportunity_columns(o, opp_by_key),
                }
            )
        for key in t["unconfirmed_keys"]:
            o = opp_by_key.get(key)
            if o:
                relation_rows.append({**token_cols, "link": "unconfirmed", "depth": None, "path": None, "status": "unconfirmed", **_opportunity_columns(o, opp_by_key)})

        for dimension in ("network", "category", "platform"):
            for key, counts in t[f"by_{dimension}"].items():
                breakdown_rows.append(
                    {"entity": "token", "entity_id": t["id"], "entity_label": t["symbol"], "ticker": t["ticker"], "dimension": dimension, "key": key, **counts}
                )

    # ---- issuers ----
    issuer_rows = []
    for i in export["issuers"]:
        row = _pick(i, ["id", "name", "cmc_label", "website", "metadata_status", "roster_size"])
        row.update(_pick(i, ["basket_underlyings_count", "basket_tokens_count", "tokens_in_defi_count"]))
        row["tokens_in_defi_share_percent"] = round(i["tokens_in_defi_count"] / i["basket_tokens_count"] * 100, 1) if i.get("basket_tokens_count") else None
        row.update(_pick(i, ["deployments_count", "supported_deployments_count", "unsupported_deployments_count", "issuer_usefulness_score"]))
        row.update(_pick(i, COUNT_FIELDS))
        row.update(_pick(i, ["basket_relative_score", "compared_underlyings_count", "mean_opportunities_per_underlying", "mean_opportunities_per_token"]))
        row.update(_pick(i, YIELD_FIELDS))
        row["best_reported_yield_keys"] = _join(i["best_reported_yield_keys"])
        row.update(_pick(i, ["direct_pools_liquidity_usd", "basket_market_cap_usd"]))
        row["eligible_platforms"] = _join(i.get("eligible_platforms"))
        row["eligible_networks"] = _join(k for k, v in i["by_network"].items() if v["eligible"])
        row["underlyings"] = _join(underlying_by_id[uid]["ticker"] for uid in i["underlying_ids"] if uid in underlying_by_id)
        row["tokens"] = _join(token_by_id[tid]["symbol"] for tid in i["token_ids"] if tid in token_by_id)
        row["coverage"] = i["coverage"]
        row["oldest_fetched_at"] = i["oldest_fetched_at"]
        row["logo"] = i["logo"]
        issuer_rows.append(row)
        for dimension in ("network", "category"):
            for key, counts in i[f"by_{dimension}"].items():
                breakdown_rows.append({"entity": "issuer", "entity_id": i["id"], "entity_label": i["name"], "ticker": None, "dimension": dimension, "key": key, **counts})
        for ticker, eligible in i["by_underlying"].items():
            breakdown_rows.append({"entity": "issuer", "entity_id": i["id"], "entity_label": i["name"], "ticker": ticker, "dimension": "underlying", "key": ticker, "raw": None, "eligible": eligible})

    # ---- opportunities: one row per position plus which basket tokens use it ----
    tokens_by_opp: dict[str, list[dict]] = {}
    for t in basket:
        for r in t["relations"]:
            tokens_by_opp.setdefault(r["key"], []).append(t)
    opportunity_rows = []
    for o in export["opportunities"]:
        users = tokens_by_opp.get(o["key"], [])
        row = _opportunity_columns(o, opp_by_key)
        row["linked_tokens_count"] = len(users)
        row["linked_tokens"] = _join(t["symbol"] for t in users)
        row["linked_tickers"] = _join(sorted({t["ticker"] for t in users}))
        row["linked_issuers"] = _join(sorted({t["issuer_name"] for t in users if t["issuer_name"]}))
        opportunity_rows.append(row)

    def table(name, rows, description, first=None):
        columns = list(first or [])
        for r in rows:
            columns += [c for c in r if c not in columns]
        return Table(name, columns, rows, description)

    return [
        table("summary", summary_rows, "Basket totals, settings and build provenance", ["metric", "value"]),
        table("underlyings", underlying_rows, "One row per underlying stock/ETF of the basket"),
        table("tokens", token_rows, "One row per tokenized asset (CMC token); kind/excluded_reason mark derivatives and other share classes"),
        table("issuers", issuer_rows, "Issuer rating over the basket"),
        table("opportunities", opportunity_rows, "One row per DeFi position (Portals) linked to at least one token"),
        table("relations", relation_rows, "Token x position links: direct, indirect (via LP/vault) and unconfirmed symbol matches"),
        table("deployments", deployment_rows, "Token contracts per network and their Portals status"),
        table("breakdowns", breakdown_rows, "Raw/eligible position counts of tokens and issuers by network, category, platform, underlying"),
    ]


def _opportunity_columns(o: dict, opp_by_key: dict) -> dict:
    row = {"opportunity_key" if f == "key" else f"position_{f}" if f in ("symbol", "name") else f: o.get(f) for f in OPPORTUNITY_FIELDS}
    row["flags"] = _join(o["flags"])
    row["tags"] = _join(o["tags"])
    row["components"] = _join(c["symbol"] or c["key"] for c in o["components"])
    row["components_count"] = len(o["components"])
    return row


# ---------- writers ----------


def _is_numeric_column(name: str) -> bool:
    return name.endswith(NUMERIC_SUFFIXES) or name in {"depth", "decimals", "value", "raw", "eligible", "rank_in_underlying", "roster_size", "cmc_rwa_id", "id", "token_id", "entity_id"}


def _cell(column: str, value):
    if value is None or value == "":
        return None
    if isinstance(value, (bool, int, float)):
        return value
    if column in TIME_FIELDS and isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)  # Excel has no tz; all times are UTC
        except ValueError:
            return value
    if _is_numeric_column(column):
        return _as_number(value)
    return value


def write_csv(tables: list[Table], out_dir: Path) -> list[Path]:
    """One `<table>.csv` per table (UTF-8, header row). Values are written as in the export: decimals as plain numbers."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for t in tables:
        path = out_dir / f"{t.name}.csv"
        fd, tmp = tempfile.mkstemp(dir=out_dir, prefix=f".{path.name}.", suffix=".tmp")
        try:
            with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=t.columns, extrasaction="ignore")
                writer.writeheader()
                for row in t.rows:
                    writer.writerow({c: _csv_value(row.get(c)) for c in t.columns})
            os.chmod(tmp, 0o644)
            os.replace(tmp, path)
        except BaseException:
            Path(tmp).unlink(missing_ok=True)
            raise
        paths.append(path)
    return paths


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def write_xlsx(tables: list[Table], path: Path) -> Path:
    """One workbook, one sheet per table: frozen header, autofilter, number formats, a README sheet listing the tables."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["sheet", "rows", "description"])
    for t in tables:
        readme.append([t.name, len(t.rows), t.description])
    readme.append([])
    readme.append(["Money, liquidity and volume are USD; APY and scores in percent; empty cell = not reported (never zero); times are UTC."])
    for cell in readme[1]:
        cell.font = Font(bold=True)
    readme.column_dimensions["A"].width = 16
    readme.column_dimensions["C"].width = 110

    for t in tables:
        ws = wb.create_sheet(t.name[:31])
        ws.append(t.columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in t.rows:
            ws.append([_cell(c, row.get(c)) for c in t.columns])
        ws.freeze_panes = "A2"
        if t.rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(t.columns))}{len(t.rows) + 1}"
        for idx, column in enumerate(t.columns, 1):
            letter = get_column_letter(idx)
            sample = [str(r.get(column)) for r in t.rows[:200] if r.get(column) is not None]
            width = max([len(column)] + [len(s) for s in sample]) if sample else len(column)
            ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
            if column in TIME_FIELDS:
                fmt = "yyyy-mm-dd hh:mm"
            elif column.endswith("_usd"):
                fmt = "#,##0.00"
            elif column.endswith("_percent"):
                fmt = "0.00"
            else:
                continue
            for cell in ws[letter][1:]:
                cell.number_format = fmt

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        os.chmod(tmp, 0o644)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
    return path


def write_tables(export: dict, fmt: str, out: Path | None, data_dir: Path) -> list[Path]:
    """`fmt`: csv (directory of files), xlsx (one workbook) or all. `out` overrides the default under data/exports/."""
    tables = build_tables(export)
    paths: list[Path] = []
    if fmt in ("csv", "all"):
        csv_dir = out if (out and fmt == "csv") else data_dir / "exports" / "csv"
        paths += write_csv(tables, csv_dir)
    if fmt in ("xlsx", "all"):
        xlsx_path = out if (out and fmt == "xlsx") else data_dir / "exports" / "rwa.xlsx"
        paths.append(write_xlsx(tables, xlsx_path))
    return paths
